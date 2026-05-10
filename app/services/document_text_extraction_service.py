from dataclasses import dataclass, field
from io import BytesIO
import re
from typing import Any

import httpx
from docx import Document
from pypdf import PdfReader

from app.core.config import get_settings

LINE_Y_TOLERANCE = 3.0
COLUMN_SPLIT_RATIO = 0.38
COLUMN_SEGMENT_GAP_THRESHOLD = 36.0
OCR_RENDER_ZOOM = 3.0
URL_RE = re.compile(r"https?://[^\s)>,;]+", re.IGNORECASE)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?:\+?\d[\d\s().-]{7,})")
BAD_GLYPH_MARKERS = ("·", "ï", "¿", "ˇ", "�")
IMPORTANT_PDF_SECTION_NAMES = {
    "experience": {"experience", "work experience"},
}
SECTION_STOP_NAMES = {
    "academic",
    "achievements",
    "awards",
    "certification",
    "certifications",
    "contact",
    "education",
    "honors",
    "honours",
    "languages",
    "profile",
    "projects",
    "skills",
    "summary",
    "technical skills",
}


@dataclass
class PdfExtractionQuality:
    has_email: bool = False
    has_phone: bool = False
    has_github: bool = False
    has_linkedin: bool = False
    has_education: bool = False
    has_projects: bool = False
    has_experience: bool = False
    corrupted_glyph_ratio: float = 0.0
    has_bad_glyphs: bool = False
    needs_header_ocr: bool = False
    needs_bottom_ocr: bool = False
    needs_full_page_ocr: bool = False


@dataclass
class OcrLine:
    text: str
    page: int
    region: str
    x0: float | None = None
    y0: float | None = None
    x1: float | None = None
    y1: float | None = None
    confidence: float | None = None


@dataclass
class OcrExtractionResult:
    raw_text: str = ""
    lines: list[OcrLine] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


class DocumentTextExtractionError(Exception):
    pass


class DocumentTextExtractionResult:
    def __init__(self, raw_text: str, method: str, warnings: list[str] | None = None):
        self.raw_text = raw_text
        self.method = method
        self.warnings = warnings or []


class DocumentTextExtractionService:
    def extract_from_signed_url(
        self,
        signed_url: str,
        file_type: str,
        file_name: str | None = None,
    ) -> DocumentTextExtractionResult:
        document_bytes = self._download_document(signed_url)
        normalized_file_type = self._normalize_file_type(file_type, file_name)

        if normalized_file_type == "PDF":
            return self._extract_pdf_text(document_bytes)

        if normalized_file_type == "DOCX":
            return self._extract_docx_text(document_bytes)

        raise DocumentTextExtractionError(
            f"Unsupported resume file type: {file_type or file_name or 'unknown'}",
        )

    def _download_document(self, signed_url: str) -> bytes:
        try:
            with httpx.Client(timeout=30.0, follow_redirects=True) as client:
                response = client.get(signed_url)
                response.raise_for_status()
                return response.content
        except httpx.HTTPError as exc:
            raise DocumentTextExtractionError(f"Failed to download resume document: {exc}") from exc

    def _normalize_file_type(self, file_type: str, file_name: str | None) -> str:
        upper_file_type = (file_type or "").upper()

        if upper_file_type in {"PDF", "DOCX"}:
            return upper_file_type

        lower_file_name = (file_name or "").lower()

        if lower_file_name.endswith(".pdf"):
            return "PDF"

        if lower_file_name.endswith(".docx"):
            return "DOCX"

        return upper_file_type

    def _extract_pdf_text(self, document_bytes: bytes) -> DocumentTextExtractionResult:
        warnings: list[str] = []

        pymupdf_text = self._try_extract_pdf_text_with_pymupdf(document_bytes, warnings)
        pypdf_text = self._try_extract_pdf_text_with_pypdf(document_bytes, warnings)

        if pymupdf_text:
            merged_text = self._merge_supplemental_pdf_text(pymupdf_text, pypdf_text, warnings)
            quality = self._analyze_pdf_extraction_quality(merged_text)
            merged_text = self._maybe_merge_local_ocr_text(document_bytes, merged_text, quality, warnings)

            return DocumentTextExtractionResult(
                raw_text=merged_text,
                method="PDF_HYBRID_PYMUPDF_PADDLEOCR" if any("PaddleOCR" in warning for warning in warnings) else "PDF_TEXT_PYMUPDF_WORDS",
                warnings=warnings,
            )

        if pypdf_text:
            quality = self._analyze_pdf_extraction_quality(pypdf_text)
            merged_text = self._maybe_merge_local_ocr_text(document_bytes, pypdf_text, quality, warnings)
            return DocumentTextExtractionResult(
                raw_text=merged_text,
                method="PDF_HYBRID_PYPDF_PADDLEOCR" if any("PaddleOCR" in warning for warning in warnings) else "PDF_TEXT_PYPDF",
                warnings=warnings,
            )

        quality = PdfExtractionQuality(needs_full_page_ocr=True, needs_header_ocr=True, needs_bottom_ocr=True)
        ocr_text = self._maybe_merge_local_ocr_text(document_bytes, "", quality, warnings)
        if ocr_text:
            return DocumentTextExtractionResult(
                raw_text=ocr_text,
                method="PDF_OCR_PADDLEOCR",
                warnings=warnings,
            )

        warnings.append(
            "No selectable text was extracted from the PDF and OCR fallback did not produce text.",
        )
        return DocumentTextExtractionResult(
            raw_text="",
            method="PDF_TEXT_EMPTY",
            warnings=warnings,
        )

    def _analyze_pdf_extraction_quality(self, text: str) -> PdfExtractionQuality:
        lowered = (text or "").lower()
        bad_glyph_count = sum((text or "").count(marker) for marker in BAD_GLYPH_MARKERS)
        corrupted_glyph_ratio = bad_glyph_count / max(len(text or ""), 1)
        has_bad_glyphs = bad_glyph_count > 0

        quality = PdfExtractionQuality(
            has_email=bool(EMAIL_RE.search(text or "")),
            has_phone=bool(PHONE_RE.search(text or "")),
            has_github="github.com" in lowered,
            has_linkedin="linkedin.com" in lowered,
            has_education="education" in lowered or "academic" in lowered,
            has_projects="projects" in lowered or "project" in lowered,
            has_experience="experience" in lowered or "work experience" in lowered,
            corrupted_glyph_ratio=corrupted_glyph_ratio,
            has_bad_glyphs=has_bad_glyphs,
        )

        quality.needs_header_ocr = has_bad_glyphs or not quality.has_email or not quality.has_phone
        quality.needs_bottom_ocr = not quality.has_experience
        quality.needs_full_page_ocr = len(text or "") < 300 or corrupted_glyph_ratio > 0.03
        return quality

    def _maybe_merge_local_ocr_text(
        self,
        document_bytes: bytes,
        primary_text: str,
        quality: PdfExtractionQuality,
        warnings: list[str],
    ) -> str:
        settings = get_settings()
        if not settings.enable_pdf_ocr_fallback:
            return self._normalize_extracted_text(primary_text)

        regions: list[str] = []
        if quality.needs_header_ocr:
            regions.extend(["header", "left_sidebar", "top_left", "top_right"])
        if quality.needs_bottom_ocr:
            regions.extend(["bottom", "bottom_right"])
        if quality.needs_full_page_ocr or settings.pdf_ocr_mode.lower() == "full_page":
            regions.append("full_page")

        if not regions:
            return self._normalize_extracted_text(primary_text)

        ocr_result = self._try_extract_pdf_text_with_paddleocr(document_bytes, regions, warnings)
        if not ocr_result.raw_text:
            return self._normalize_extracted_text(primary_text)

        merged_text = self._merge_ocr_result(primary_text, ocr_result, warnings)
        return self._normalize_extracted_text(merged_text)

    def _try_extract_pdf_text_with_paddleocr(
        self,
        document_bytes: bytes,
        regions: list[str],
        warnings: list[str],
    ) -> OcrExtractionResult:
        try:
            import fitz  # PyMuPDF
            from paddleocr import PaddleOCR
            from PIL import Image
        except ImportError as exc:
            warnings.append(f"PaddleOCR fallback skipped because optional OCR dependencies are missing: {exc}")
            return OcrExtractionResult()

        settings = get_settings()
        max_pages = max(settings.pdf_ocr_max_pages, 1)
        min_confidence = settings.pdf_ocr_min_confidence
        lang = (settings.pdf_ocr_languages or "en").split("+")[0]

        try:
            ocr = PaddleOCR(use_angle_cls=True, lang=lang, show_log=False)
        except Exception as exc:
            warnings.append(f"PaddleOCR fallback initialization failed: {exc}")
            return OcrExtractionResult()

        ocr_lines: list[OcrLine] = []
        unique_regions = list(dict.fromkeys(regions))

        try:
            with fitz.open(stream=document_bytes, filetype="pdf") as document:
                for page_index, page in enumerate(document):
                    if page_index >= max_pages:
                        break
                    for region_name, rect in self._ocr_regions_for_page(page, unique_regions):
                        pix = page.get_pixmap(
                            matrix=fitz.Matrix(OCR_RENDER_ZOOM, OCR_RENDER_ZOOM),
                            clip=rect,
                            alpha=False,
                        )
                        image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                        result = ocr.ocr(image, cls=True)
                        ocr_lines.extend(
                            self._paddleocr_result_to_lines(
                                result,
                                page_index + 1,
                                region_name,
                                rect,
                                OCR_RENDER_ZOOM,
                                min_confidence,
                            )
                        )
        except Exception as exc:
            warnings.append(f"PaddleOCR fallback failed: {exc}")
            return OcrExtractionResult(lines=ocr_lines)

        sorted_lines = sorted(ocr_lines, key=lambda line: (line.page, line.y0 or 0, line.x0 or 0))
        raw_text = "\n".join(line.text for line in sorted_lines if line.text)
        if raw_text:
            warnings.append(f"PaddleOCR fallback extracted text from regions: {', '.join(unique_regions)}")

        return OcrExtractionResult(raw_text=raw_text, lines=sorted_lines)

    def _ocr_regions_for_page(self, page: Any, regions: list[str]) -> list[tuple[str, Any]]:
        import fitz

        width = float(page.rect.width)
        height = float(page.rect.height)
        mapping = {
            "header": fitz.Rect(0, 0, width, height * 0.28),
            "left_sidebar": fitz.Rect(0, 0, width * 0.42, height),
            "top_left": fitz.Rect(0, 0, width * 0.45, height * 0.45),
            "top_right": fitz.Rect(width * 0.35, 0, width, height * 0.35),
            "bottom": fitz.Rect(0, height * 0.62, width, height),
            "bottom_right": fitz.Rect(width * 0.35, height * 0.55, width, height),
            "full_page": fitz.Rect(0, 0, width, height),
        }
        return [(region, mapping[region]) for region in regions if region in mapping]

    def _paddleocr_result_to_lines(
        self,
        result: Any,
        page: int,
        region: str,
        rect: Any,
        zoom: float,
        min_confidence: float,
    ) -> list[OcrLine]:
        lines: list[OcrLine] = []
        if not result:
            return lines

        for page_result in result:
            for item in page_result or []:
                if not item or len(item) < 2:
                    continue
                box, recognition = item[0], item[1]
                if not recognition or len(recognition) < 2:
                    continue
                text, confidence = recognition[0], float(recognition[1] or 0)
                text = str(text).strip()
                if not text or confidence < min_confidence:
                    continue

                try:
                    x_values = [float(point[0]) / zoom + float(rect.x0) for point in box]
                    y_values = [float(point[1]) / zoom + float(rect.y0) for point in box]
                except Exception:
                    x_values = []
                    y_values = []

                lines.append(
                    OcrLine(
                        text=text,
                        page=page,
                        region=region,
                        x0=min(x_values) if x_values else None,
                        y0=min(y_values) if y_values else None,
                        x1=max(x_values) if x_values else None,
                        y1=max(y_values) if y_values else None,
                        confidence=confidence,
                    )
                )
        return lines

    def _merge_ocr_result(self, primary_text: str, ocr_result: OcrExtractionResult, warnings: list[str]) -> str:
        primary = self._normalize_extracted_text(primary_text)
        ocr_text = self._normalize_extracted_text(ocr_result.raw_text)
        if not primary:
            return ocr_text

        blocks: list[str] = [primary]
        header_text = self._ocr_text_for_regions(ocr_result, {"header", "left_sidebar", "top_left", "top_right"})
        bottom_text = self._ocr_text_for_regions(ocr_result, {"bottom", "bottom_right", "full_page"})

        if self._has_corrupted_glyphs(primary) and header_text:
            blocks.insert(0, f"[OCR_HEADER]\n{header_text}\n[/OCR_HEADER]")
            warnings.append("Merged OCR header text to recover corrupted personal information")

        if not self._contains_any_section(primary, IMPORTANT_PDF_SECTION_NAMES["experience"]):
            experience_section = self._extract_named_section(bottom_text, IMPORTANT_PDF_SECTION_NAMES["experience"])
            if experience_section:
                blocks.append(f"[OCR_RECOVERED_EXPERIENCE]\n{experience_section}")
                warnings.append("Merged OCR bottom text to recover missing experience section")

        if len(blocks) == 1 and ocr_text:
            blocks.append(f"[OCR_SUPPLEMENTAL]\n{ocr_text}")
            warnings.append("Merged OCR supplemental text because no targeted OCR block matched")

        return "\n".join(blocks)

    def _ocr_text_for_regions(self, ocr_result: OcrExtractionResult, regions: set[str]) -> str:
        lines = [line for line in ocr_result.lines if line.region in regions]
        lines = sorted(lines, key=lambda line: (line.page, line.y0 or 0, line.x0 or 0))
        return self._normalize_extracted_text("\n".join(line.text for line in lines))

    def _has_corrupted_glyphs(self, text: str) -> bool:
        return any(marker in (text or "") for marker in BAD_GLYPH_MARKERS)

    def _contains_any_section(self, text: str, aliases: set[str]) -> bool:
        normalized = self._normalize_section_key(text)
        return any(alias in normalized for alias in aliases)

    def _merge_supplemental_pdf_text(self, primary_text: str, fallback_text: str, warnings: list[str]) -> str:
        """Append important sections that PyMuPDF layout extraction missed.

        PyMuPDF words preserve layout better for multi-column resumes, so it
        remains the primary source. Some PDFs, however, drop bottom sections or
        text boxes in the word stream. pypdf sometimes still exposes those
        sections. This merger only appends known important missing sections,
        avoiding a broad raw-text concatenation that would duplicate most of the
        resume or destroy layout ordering.
        """
        primary = self._normalize_extracted_text(primary_text)
        fallback = self._normalize_extracted_text(fallback_text)

        if not primary or not fallback:
            return primary or fallback

        supplemental_sections: list[str] = []

        for section_name, aliases in IMPORTANT_PDF_SECTION_NAMES.items():
            if self._contains_any_section(primary, aliases):
                continue

            section_text = self._extract_named_section(fallback, aliases)
            if section_text:
                supplemental_sections.append(section_text)
                warnings.append(f"Merged missing PDF section from pypdf fallback: {section_name}")

        if not supplemental_sections:
            return primary

        return self._normalize_extracted_text("\n".join([primary, *supplemental_sections]))

    def _extract_named_section(self, text: str, aliases: set[str]) -> str:
        lines = [line.strip() for line in (text or "").splitlines() if line.strip()]
        collected: list[str] = []
        collecting = False

        for line in lines:
            normalized = self._normalize_section_key(line)

            if not collecting and normalized in aliases:
                collecting = True
                collected.append(line)
                continue

            if collecting and normalized in SECTION_STOP_NAMES and normalized not in aliases:
                break

            if collecting:
                collected.append(line)

        return "\n".join(collected).strip()

    def _normalize_section_key(self, value: str) -> str:
        lowered = (value or "").lower()
        lowered = re.sub(r"[^a-z0-9+#./ ]+", " ", lowered)
        return re.sub(r"\s+", " ", lowered).strip()

    def _try_extract_pdf_text_with_pymupdf(
        self,
        document_bytes: bytes,
        warnings: list[str],
    ) -> str:
        try:
            import fitz  # PyMuPDF
        except ImportError:
            warnings.append("PyMuPDF is not installed; falling back to pypdf for PDF text extraction.")
            return ""

        try:
            with fitz.open(stream=document_bytes, filetype="pdf") as document:
                page_texts = [self._extract_pymupdf_page_text(page) for page in document]
        except Exception as exc:  # PyMuPDF raises multiple parser-specific exceptions.
            warnings.append(f"PyMuPDF PDF extraction failed; falling back to pypdf: {exc}")
            return ""

        return self._normalize_extracted_text("\n".join(page_texts))

    def _extract_pymupdf_page_text(self, page: Any) -> str:
        words = page.get_text("words", sort=False)
        if not words:
            text = page.get_text("text", sort=True)
            return self._append_page_links(text, self._extract_page_links(page))

        page_width = float(page.rect.width)
        page_links = self._extract_page_links(page)
        normalized_words = []

        for word in words:
            if len(word) < 5:
                continue

            x0, y0, x1, y1, text = word[:5]
            item = {
                "x0": float(x0),
                "y0": float(y0),
                "x1": float(x1),
                "y1": float(y1),
                "text": str(text).strip(),
            }
            if item["text"]:
                normalized_words.append(item)

        line_segments = self._words_to_line_segments(normalized_words, page_links, page_width)
        page_text = self._segments_to_reading_order_text(line_segments, page_width)

        return self._append_unmatched_links(page_text, page_links)

    def _extract_page_links(self, page: Any) -> list[dict[str, Any]]:
        links: list[dict[str, Any]] = []

        try:
            raw_links = page.get_links()
        except Exception:
            return links

        for raw_link in raw_links:
            uri = self._normalize_url(raw_link.get("uri", ""))
            rect = raw_link.get("from")
            if not uri or rect is None:
                continue

            links.append(
                {
                    "url": uri,
                    "x0": float(rect.x0),
                    "y0": float(rect.y0),
                    "x1": float(rect.x1),
                    "y1": float(rect.y1),
                    "matched": False,
                },
            )

        return links

    def _words_to_line_segments(
        self,
        words: list[dict[str, Any]],
        links: list[dict[str, Any]],
        page_width: float,
    ) -> list[dict[str, Any]]:
        if not words:
            return []

        segments: list[dict[str, Any]] = []
        gap_threshold = max(COLUMN_SEGMENT_GAP_THRESHOLD, page_width * 0.06)

        for line in self._group_items_into_lines(words):
            ordered_line = sorted(line, key=lambda word: float(word["x0"]))
            current_segment: list[dict[str, Any]] = []

            for word in ordered_line:
                if current_segment:
                    previous_word = current_segment[-1]
                    gap = float(word["x0"]) - float(previous_word["x1"])
                    if gap > gap_threshold:
                        segments.append(self._build_text_segment(current_segment, links))
                        current_segment = []

                current_segment.append(word)

            if current_segment:
                segments.append(self._build_text_segment(current_segment, links))

        return segments

    def _build_text_segment(
        self,
        words: list[dict[str, Any]],
        links: list[dict[str, Any]],
    ) -> dict[str, Any]:
        ordered_words = sorted(words, key=lambda word: float(word["x0"]))
        return {
            "x0": min(float(word["x0"]) for word in ordered_words),
            "y0": min(float(word["y0"]) for word in ordered_words),
            "x1": max(float(word["x1"]) for word in ordered_words),
            "y1": max(float(word["y1"]) for word in ordered_words),
            "text": " ".join(str(word["text"]) for word in ordered_words).strip(),
            "urls": self._urls_for_line(ordered_words, links),
        }

    def _segments_to_reading_order_text(
        self,
        segments: list[dict[str, Any]],
        page_width: float,
    ) -> str:
        if not segments:
            return ""

        column_start_y = self._detect_two_column_start_y(segments, page_width)

        if column_start_y is None:
            ordered_segments = sorted(segments, key=lambda segment: (float(segment["y0"]), float(segment["x0"])))
        else:
            before_columns = [
                segment
                for segment in segments
                if float(segment["y0"]) < column_start_y - LINE_Y_TOLERANCE
            ]
            column_segments = [
                segment
                for segment in segments
                if float(segment["y0"]) >= column_start_y - LINE_Y_TOLERANCE
            ]
            split_x = page_width * COLUMN_SPLIT_RATIO
            left_column = [segment for segment in column_segments if float(segment["x0"]) < split_x]
            right_column = [segment for segment in column_segments if float(segment["x0"]) >= split_x]

            ordered_segments = [
                *sorted(before_columns, key=lambda segment: (float(segment["y0"]), float(segment["x0"]))),
                *sorted(left_column, key=lambda segment: (float(segment["y0"]), float(segment["x0"]))),
                *sorted(right_column, key=lambda segment: (float(segment["y0"]), float(segment["x0"]))),
            ]

        output_lines: list[str] = []
        for segment in ordered_segments:
            output_lines.extend(self._segment_output_lines(segment))

        return "\n".join(output_lines)

    def _detect_two_column_start_y(
        self,
        segments: list[dict[str, Any]],
        page_width: float,
    ) -> float | None:
        split_x = page_width * COLUMN_SPLIT_RATIO

        for line in self._group_items_into_lines(segments):
            if len(line) < 2:
                continue

            has_left_segment = any(float(segment["x0"]) < split_x for segment in line)
            has_right_segment = any(float(segment["x0"]) >= split_x for segment in line)

            if has_left_segment and has_right_segment:
                return min(float(segment["y0"]) for segment in line)

        return None

    def _segment_output_lines(self, segment: dict[str, Any]) -> list[str]:
        lines: list[str] = []
        text = str(segment.get("text") or "").strip()
        if text:
            lines.append(text)

        for url in segment.get("urls") or []:
            lines.append(str(url))

        return lines

    def _group_items_into_lines(self, items: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
        sorted_items = sorted(items, key=lambda item: (float(item["y0"]), float(item["x0"])))
        lines: list[list[dict[str, Any]]] = []

        for item in sorted_items:
            if not lines:
                lines.append([item])
                continue

            current_line = lines[-1]
            current_y = sum(float(existing_item["y0"]) for existing_item in current_line) / len(current_line)
            if abs(float(item["y0"]) - current_y) <= LINE_Y_TOLERANCE:
                current_line.append(item)
            else:
                lines.append([item])

        return lines

    def _urls_for_line(
        self,
        line: list[dict[str, Any]],
        links: list[dict[str, Any]],
    ) -> list[str]:
        if not links or not line:
            return []

        min_x = min(float(word["x0"]) for word in line)
        max_x = max(float(word["x1"]) for word in line)
        min_y = min(float(word["y0"]) for word in line)
        max_y = max(float(word["y1"]) for word in line)
        urls: list[str] = []

        for link in links:
            if link["matched"]:
                continue
            horizontally_overlaps = float(link["x1"]) >= min_x - 4 and float(link["x0"]) <= max_x + 4
            vertically_overlaps = float(link["y1"]) >= min_y - 4 and float(link["y0"]) <= max_y + 4

            if horizontally_overlaps and vertically_overlaps:
                link["matched"] = True
                urls.append(str(link["url"]))

        return self._unique_urls(urls)

    def _append_page_links(self, text: str, links: list[dict[str, Any]]) -> str:
        urls = self._unique_urls([str(link["url"]) for link in links])
        if not urls:
            return text

        return "\n".join([text, *urls])

    def _append_unmatched_links(self, text: str, links: list[dict[str, Any]]) -> str:
        urls = self._unique_urls([str(link["url"]) for link in links if not link["matched"]])
        if not urls:
            return text

        existing_urls = set(URL_RE.findall(text or ""))
        missing_urls = [url for url in urls if url not in existing_urls]
        if not missing_urls:
            return text

        return "\n".join([text, *missing_urls])

    def _unique_urls(self, urls: list[str]) -> list[str]:
        result = []
        seen = set()

        for url in urls:
            normalized_url = self._normalize_url(url)
            if not normalized_url:
                continue
            key = normalized_url.rstrip("/").lower()
            if key in seen:
                continue
            seen.add(key)
            result.append(normalized_url)

        return result

    def _normalize_url(self, value: str) -> str | None:
        cleaned = (value or "").strip().rstrip(".,;)")
        if not re.match(r"^https?://", cleaned, flags=re.IGNORECASE):
            return None
        return cleaned

    def _try_extract_pdf_text_with_pypdf(
        self,
        document_bytes: bytes,
        warnings: list[str],
    ) -> str:
        try:
            reader = PdfReader(BytesIO(document_bytes))
            page_texts = [(page.extract_text() or "") for page in reader.pages]
        except Exception as exc:  # pypdf raises multiple parser-specific exceptions.
            warnings.append(f"pypdf PDF extraction failed: {exc}")
            return ""

        return self._normalize_extracted_text("\n".join(page_texts))

    def _extract_docx_text(self, document_bytes: bytes) -> DocumentTextExtractionResult:
        try:
            document = Document(BytesIO(document_bytes))
            raw_text_parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        raw_text_parts.append(row_text)
            raw_text = "\n".join(raw_text_parts).strip()
        except Exception as exc:  # python-docx may raise zip/xml parser errors for invalid docs.
            raise DocumentTextExtractionError(f"Failed to extract DOCX text: {exc}") from exc

        raw_text = self._normalize_extracted_text(raw_text)
        warnings: list[str] = []

        if not raw_text:
            warnings.append("No text was extracted from the DOCX document.")

        return DocumentTextExtractionResult(
            raw_text=raw_text,
            method="DOCX_TEXT",
            warnings=warnings,
        )

    def _normalize_extracted_text(self, text: str) -> str:
        return "\n".join(
            re.sub(r"[ \t\f\v]+", " ", line).strip()
            for line in (text or "").replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
            if line.strip()
        ).strip()


document_text_extraction_service = DocumentTextExtractionService()
